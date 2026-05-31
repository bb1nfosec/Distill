"""Rebuild slide 2 with cleaner layout and less text."""

import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

import shutil
from pptx import Presentation
from pptx.util import Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from lxml import etree

SRC = r'C:\Users\528875\OneDrive - Genuine Parts Company\Microsoft Teams Chat Files\Distill_Token_Optimization 1.pptx'
DST = r'C:\Users\528875\Distill_Token_Optimization.pptx'
shutil.copy2(SRC, DST)

prs = Presentation(DST)
print(f"Loaded: {len(prs.slides)} slides")

# --- Helpers ---

A_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'


def make_textbox(slide, left, top, width, height, text, sz, bold, color, align="l"):
    box = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(height))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]

    # Set alignment
    al = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
    p.alignment = al.get(align, PP_ALIGN.LEFT)

    # Add run with explicit formatting (not defRPr - avoids alignment issues)
    run = p.add_run()
    run.text = text
    run.font.size = Emu(sz * 12700)  # points to EMU
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = "Segoe UI"
    return box


def make_card(slide, left, top, width, height, fill):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Emu(left), Emu(top), Emu(width), Emu(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    shape.line.fill.background()
    shape.adjustments[0] = 0.06
    return shape


def make_circle(slide, left, top, size, fill, label):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Emu(left), Emu(top), Emu(size), Emu(size))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = label
    run.font.size = Emu(16 * 12700)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string("FFFFFF")
    run.font.name = "Segoe UI"
    return shape


# --- Build new slide ---

blank_layout = None
for layout in prs.slide_layouts:
    if layout.name == "Blank":
        blank_layout = layout
        break
if not blank_layout:
    blank_layout = prs.slide_layouts[6]

slide = prs.slides.add_slide(blank_layout)

# Background
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = RGBColor(0x0F, 0x17, 0x2A)

# Top accent bar
bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Emu(12192000), Emu(54864))
bar.fill.solid()
bar.fill.fore_color.rgb = RGBColor.from_string("00D4AA")
bar.line.fill.background()

# Title + subtitle
make_textbox(slide, 731520, 365760, 10728960, 640080,
             "What Does Distill Do?", 36, True, "FFFFFF")
make_textbox(slide, 731520, 960120, 10728960, 365760,
             "Distill never modifies your source code. It controls what the LLM sees and what it costs.",
             16, False, "94A3B8")

# --- Left column ---
make_textbox(slide, 731520, 1554480, 5303520, 365760,
             "CLI Tools", 24, True, "00D4AA")

left_items = [
    ("1", "distill scan",    "Token cost per file"),
    ("2", "distill analyze", "Detect waste patterns"),
    ("3", "distill check",   "CI budget gate"),
    ("4", "distill fix",     "Generate .llmignore"),
]

card_w = 5120640
card_h = 548640
card_x = 731520
gap = 137160
y = 1965960

for num, title, desc in left_items:
    make_card(slide, card_x, y, card_w, card_h, "16213E")
    make_circle(slide, card_x + 137160, y + 114300, 320040, "00D4AA", num)
    make_textbox(slide, card_x + 548640, y + 91440, card_w - 640080, 228600,
                 title, 16, True, "FFFFFF")
    make_textbox(slide, card_x + 548640, y + 320040, card_w - 640080, 228600,
                 desc, 13, False, "94A3B8")
    y += card_h + gap

# --- Right column ---
sec2_x = 6309360

make_textbox(slide, sec2_x, 1554480, 5303520, 365760,
             "Python Adapters", 24, True, "58A6FF")

right_items = [
    ("1", "Auto-Compaction",    "Summarize history when context fills"),
    ("2", "Prompt Caching",     "90% cost reduction on repeated context"),
    ("3", "Lean Prompts",       "Shorter defaults, opt out with lean_mode=False"),
    ("4", "Lazy File Loading",  "Load on demand, head+tail truncation"),
]

y = 1965960
for num, title, desc in right_items:
    make_card(slide, sec2_x, y, card_w, card_h, "16213E")
    make_circle(slide, sec2_x + 137160, y + 114300, 320040, "58A6FF", num)
    make_textbox(slide, sec2_x + 548640, y + 91440, card_w - 640080, 228600,
                 title, 16, True, "FFFFFF")
    make_textbox(slide, sec2_x + 548640, y + 320040, card_w - 640080, 228600,
                 desc, 13, False, "94A3B8")
    y += card_h + gap

# Bottom note
make_textbox(slide, 731520, 5943600, 10728960, 365760,
             "Your code stays untouched. Only .llmignore files are written (by distill fix).",
             14, False, "58A6FF")

# --- Move to position 2 ---
sldIdLst = prs.element.sldIdLst
sldIds = list(sldIdLst)
new_sldId = sldIds[-1]
sldIdLst.remove(new_sldId)
sldIds = list(sldIdLst)
sldIds[0].addnext(new_sldId)

# --- Verify ---
print("\nFinal order:")
for i, s in enumerate(prs.slides):
    txt = ""
    for sh in s.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                t = p.text.strip()
                if t and len(t) > 10:
                    txt = t[:60]
                    break
        if txt:
            break
    print(f"  Slide {i+1}: {txt}")

prs.save(DST)
print(f"\nSaved. {len(prs.slides)} slides total.")
